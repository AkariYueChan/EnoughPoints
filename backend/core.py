"""Excel import and course planning domain logic; all external content is data."""
import io
import math
import re
import unicodedata
import uuid
import zipfile
from decimal import Decimal, InvalidOperation

import openpyxl
import xlrd

MAX_ROWS, MAX_COLS, MAX_COURSES = 10000, 100, 2000
FIELDS = ('code', 'name', 'type', 'category', 'credits', 'terms')
ALIASES = {
    'code': ['课程编号', '课程代码', '课程号', 'coursecode', 'code'],
    'name': ['课程名称', '课程名', '科目', 'coursename', 'name'],
    'type': ['课程属性', '课程性质', '必修选修', '修读性质', 'type'],
    'category': ['课程体系', '课程类别', '课程分类', '课程模块', '课程性质', 'category'],
    'credits': ['学分', '课程学分', 'credit', 'credits'],
    'terms': ['开设学期', '建议修读学期', '开课学期', '学期', 'semester', 'terms'],
}


def txt(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def norm(value):
    return re.sub(r'[\s\n/（）()_\-]', '', unicodedata.normalize('NFKC', txt(value))).lower()


def course_type(value):
    value = norm(value)
    if value in ('必修', '必修课', 'required', 'compulsory'):
        return '必修'
    if value in ('选修', '选修课', '任选', '限选', '任选课', '限选课', 'elective', 'optional'):
        return '选修'
    return txt(value)


def parse_terms(value):
    if value is None or value == '' or value == []:
        return []
    if isinstance(value, list):
        if any(isinstance(n, bool) or not isinstance(n, int) or not 1 <= n <= 20 for n in value):
            raise ValueError('学期必须是 1–20 的整数')
        return sorted(set(value))
    s = unicodedata.normalize('NFKC', txt(value))
    for y, number in [('一', 1), ('二', 2), ('三', 3), ('四', 4), ('五', 5)]:
        for half, delta in [('上', -1), ('下', 0)]:
            s = s.replace('大' + y + half, str(number * 2 + delta))
    for zh, n in [('十二',12),('十一',11),('十',10),('九',9),('八',8),('七',7),('六',6),('五',5),('四',4),('三',3),('二',2),('一',1)]:
        s = s.replace(zh, str(n))
    s = re.sub(r'第|学期', '', s)
    s = re.sub(r'[、,，;；/\s]+', ',', s.strip())
    s = re.sub(r'[~～至—–]', '-', s)
    result = []
    for part in s.strip(',').split(','):
        if re.fullmatch(r'\d{1,2}', part):
            result.append(int(part))
        elif re.fullmatch(r'\d{1,2}-\d{1,2}', part):
            a, b = map(int, part.split('-'))
            if a > b or b > 20:
                raise ValueError('学期范围无效')
            result.extend(range(a, b + 1))
        else:
            raise ValueError('无法识别学期，请填写如 1,2 或 1-3')
    if any(not 1 <= n <= 20 for n in result):
        raise ValueError('学期必须在 1–20 之间')
    return sorted(set(result))


def number(value, label='学分', maximum=100):
    if isinstance(value, bool) or txt(value) == '':
        raise ValueError(label + '不能为空')
    try:
        d = Decimal(txt(value))
    except InvalidOperation:
        raise ValueError(label + '必须为数字')
    if not d.is_finite() or d < 0 or d > maximum or d != d.quantize(Decimal('.01')):
        raise ValueError(label + f'须为 0–{maximum} 的数字，最多两位小数')
    return float(d)


def read_excel(data, filename):
    ext = filename.lower().rsplit('.', 1)[-1]
    sheets = []
    if ext == 'xlsx':
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                entries = archive.infolist()
                if len(entries) > 3000 or sum(e.file_size for e in entries) > 50 * 1024 * 1024:
                    raise ValueError('工作簿解压后过大（限制 50 MB）')
            wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
            cached = openpyxl.load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
            if len(wb.worksheets) > 30:
                raise ValueError('工作表数量超过 30')
            for sheet in wb.worksheets:
                if sheet.max_row > MAX_ROWS or sheet.max_column > MAX_COLS:
                    raise ValueError('工作表超过 10,000 行或 100 列，请裁剪后重试')
                rows, formulas = [], []
                for row in sheet:
                    values, flags = [], []
                    for cell in row:
                        is_formula = cell.data_type == 'f'
                        value = cached[sheet.title][cell.coordinate].value if is_formula else cell.value
                        # Preserve course IDs formatted as 000000 when their original value is numeric.
                        if isinstance(value, (int, float)) and re.fullmatch(r'0{2,}', cell.number_format or ''):
                            value = str(int(value)).zfill(len(cell.number_format))
                        values.append(txt(value))
                        flags.append('missing' if is_formula and value is None else 'cached' if is_formula else '')
                    rows.append(values)
                    formulas.append(flags)
                merges = [(r.min_row - 1, r.max_row, r.min_col - 1, r.max_col) for r in sheet.merged_cells.ranges]
                sheets.append(dict(name=sheet.title, rows=rows, formulas=formulas, merges=merges))
            wb.close()
            cached.close()
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError('无法读取 xlsx，文件可能损坏、加密或格式不符') from exc
    elif ext == 'xls':
        try:
            wb = xlrd.open_workbook(file_contents=data, formatting_info=True, on_demand=True)
            if wb.nsheets > 30:
                raise ValueError('工作表数量超过 30')
            for sheet in wb.sheets():
                if sheet.nrows > MAX_ROWS or sheet.ncols > MAX_COLS:
                    raise ValueError('工作表超过 10,000 行或 100 列')
                rows = []
                for row in range(sheet.nrows):
                    vals = []
                    for col in range(sheet.ncols):
                        cell = sheet.cell(row, col)
                        value = txt(cell.value)
                        if cell.ctype == xlrd.XL_CELL_NUMBER:
                            fmt = wb.format_map[wb.xf_list[cell.xf_index].format_key].format_str
                            if re.fullmatch(r'0{2,}', fmt):
                                value = str(int(cell.value)).zfill(len(fmt))
                        vals.append(value)
                    rows.append(vals)
                sheets.append(dict(name=sheet.name, rows=rows, formulas=[], merges=sheet.merged_cells))
            wb.release_resources()
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError('无法读取 xls，文件可能损坏、加密或格式不符') from exc
    else:
        raise ValueError('仅支持 .xlsx 和 .xls 文件')
    if not sheets:
        raise ValueError('文件中没有可读取的工作表')
    return sheets


def infer_mapping(rows, header):
    labels = rows[header] if header < len(rows) else []
    mapping = {}
    for field, aliases in ALIASES.items():
        candidates = [(aliases.index(norm(v)), c) for c, v in enumerate(labels) if norm(v) in aliases]
        mapping[field] = min(candidates)[1] if candidates else None
    # Some universities swap the meanings of 课程性质/课程属性. Inspect data too.
    scores = []
    for col in range(len(labels)):
        score = sum(course_type(r[col]) in ('必修', '选修') for r in rows[header + 1:header + 41] if col < len(r))
        if score:
            scores.append((score, col))
    if scores:
        mapping['type'] = max(scores)[1]
    if mapping['category'] == mapping['type']:
        mapping['category'] = None
    return mapping


def describe_sheets(sheets):
    result = []
    for s in sheets:
        best = (0, 0)
        for i, row in enumerate(s['rows'][:50]):
            score = sum(any(norm(v) in aliases for v in row) for aliases in ALIASES.values())
            if score > best[0]:
                best = (score, i)
        header = best[1]
        result.append(dict(name=s['name'], rowCount=len(s['rows']), header=header + 1,
                           mapping=infer_mapping(s['rows'], header), preview=s['rows'][:60]))
    return result


def parse_sheet(sheet, header, mapping):
    if not isinstance(header, int) or isinstance(header, bool) or not 1 <= header <= len(sheet['rows']):
        raise ValueError('表头行无效')
    cols = max((len(r) for r in sheet['rows']), default=0)
    for key in FIELDS:
        col = mapping.get(key)
        if col is not None and (isinstance(col, bool) or not isinstance(col, int) or not 0 <= col < cols):
            raise ValueError('字段列映射无效')
    if any(mapping.get(k) is None for k in ('name', 'credits', 'type')):
        raise ValueError('请指定课程名称、学分和必修/选修所在列')
    mapped = [mapping.get(k) for k in FIELDS if mapping.get(k) is not None]
    if len(mapped) != len(set(mapped)):
        raise ValueError('不同字段不能使用同一列')
    courses, skipped = [], 0
    for ri in range(header, len(sheet['rows'])):
        row = sheet['rows'][ri]
        nonempty = [v for v in row if txt(v)]
        if not nonempty or any(re.fullmatch(r'(小计|合计|总计|学分小计|学分合计|总学分)([：:\s].*)?', txt(v)) for v in row[:4]):
            skipped += 1
            continue
        def value(field):
            ci = mapping.get(field)
            if ci is None:
                return ''
            val = row[ci] if ci < len(row) else ''
            if not val and field in ('category', 'type'):
                for a, b, c, d in sheet['merges']:
                    if a <= ri < b and c <= ci < d:
                        val = sheet['rows'][a][c]
                        break
            return val
        name, code, credit = value('name'), value('code'), value('credits')
        # Skip second-tier headers (such as 学时分类) only if course fields are empty.
        raw_type_col = mapping.get('type')
        raw_type = row[raw_type_col] if raw_type_col is not None and raw_type_col < len(row) else ''
        if not name and not code and not credit and not raw_type and not value('terms'):
            skipped += 1
            continue
        if norm(name) in ALIASES['name'] and norm(credit) in ALIASES['credits']:
            skipped += 1
            continue
        category = re.sub(r'\s*[（(]\s*应修.*$', '', value('category')).strip()
        if category.startswith('公共基础课'):
            category = '公共基础课'
        try:
            terms = parse_terms(value('terms'))
        except ValueError:
            terms = value('terms')
        notes = []
        if value('type') in ('限选', '限选课', '任选', '任选课'):
            notes.append('原课程性质为“'+value('type')+'”，已归入选修；请按培养方案核对类别与最低学分要求')
        if sheet['formulas']:
            for field in FIELDS:
                ci = mapping.get(field)
                if ci is not None and sheet['formulas'][ri][ci]:
                    notes.append(f'{field} 列使用公式' + ('且无缓存，请在 Excel 计算并保存或手动修正' if sheet['formulas'][ri][ci] == 'missing' else '，使用上次保存的结果'))
        courses.append(dict(id=str(uuid.uuid4()), code=code, name=name, type=course_type(value('type')),
                            category=category or '未分类', credits=credit, terms=terms,
                            source={'sheet': sheet['name'], 'row': ri + 1}, notes=notes))
        if len(courses) > MAX_COURSES:
            raise ValueError('课程超过 2,000 门，请拆分后导入')
    clean, issues = validate_courses(courses)
    return dict(courses=clean, issues=issues, skipped=skipped)


def validate_courses(courses):
    if not isinstance(courses, list) or not 1 <= len(courses) <= MAX_COURSES:
        raise ValueError('课程数量须为 1–2,000 门')
    clean, issues, seen_ids, seen_codes, seen_names = [], [], set(), set(), set()
    for index, original in enumerate(courses):
        if not isinstance(original, dict):
            raise ValueError('课程数据结构无效')
        c = {k: original.get(k, '') for k in ('id','code','name','type','category','credits','terms','source','notes')}
        c['id'] = txt(c['id']) or str(uuid.uuid4())
        c['code'], c['name'], c['category'] = txt(c['code']), txt(c['name']), txt(c['category']) or '未分类'
        c['type'] = course_type(c['type'])
        def issue(level, field, message):
            issues.append(dict(level=level, index=index, id=c['id'], field=field, message=message))
        if c['id'] in seen_ids:
            issue('error', 'id', '内部标识重复，请删除重复记录重新添加')
        seen_ids.add(c['id'])
        if not c['name']:
            issue('error', 'name', '课程名称不能为空')
        if any(len(txt(c[k])) > 300 for k in ('id','code','name','category','type')):
            issue('error', 'name', '文本字段不能超过 300 字符')
        if c['type'] not in ('必修', '选修'):
            issue('error', 'type', '请选择必修或选修')
        try:
            c['credits'] = number(c['credits'])
        except ValueError as exc:
            issue('error', 'credits', str(exc))
        try:
            c['terms'] = parse_terms(c['terms'])
            if not c['terms']:
                issue('warning', 'terms', '未指定学期，将列入待安排')
        except ValueError as exc:
            issue('error', 'terms', str(exc))
        if c['code'] and c['code'] in seen_codes:
            issue('warning', 'code', '课程编号重复，请确认是否应删除或合并；保留时按独立课程计分')
        elif not c['code']:
            issue('warning', 'code', '未提供课程编号，使用内部标识保存')
        if c['name'] in seen_names:
            issue('warning', 'name', '课程名称重复，请核对是否为不同课程')
        if c['category'] not in ('公共基础课', '专业课'):
            issue('error', 'category', '课程类别只能是公共基础课或专业课，请手动选择对应类别')
        for note in c['notes'] if isinstance(c['notes'], list) else []:
            issue('warning', 'source', txt(note))
        seen_codes.add(c['code'])
        seen_names.add(c['name'])
        clean.append(c)
    return clean, issues


def validate_plan(payload):
    courses, issues = validate_courses(payload.get('courses'))
    name = txt(payload.get('name'))
    if not name or len(name) > 120:
        raise ValueError('方案名称须为 1–120 字符')
    count = payload.get('termCount', 6)
    if isinstance(count, bool) or not isinstance(count, int) or not 1 <= count <= 20:
        raise ValueError('学期数须为 1–20 的整数')
    for i, c in enumerate(courses):
        if isinstance(c['terms'], list) and any(t > count for t in c['terms']):
            issues.append(dict(level='error', index=i, id=c['id'], field='terms', message='开设学期超出方案学期数'))
    raw = payload.get('goals') or {}
    if not isinstance(raw, dict):
        raise ValueError('学分目标格式无效')
    goals = {}
    for key in ('total', 'elective'):
        goals[key] = None if raw.get(key) in (None, '') else number(raw[key], '目标学分', 10000)
    cats = raw.get('categories') or {}
    if not isinstance(cats, dict) or len(cats) > 100:
        raise ValueError('类别目标格式无效')
    goals['categories'] = {txt(k): number(v, '类别目标', 10000) for k, v in cats.items() if v not in ('', None)}
    if any(k not in {c['category'] for c in courses} for k in goals['categories']):
        raise ValueError('目标中包含不存在的课程类别，请重新核对')
    if not any(i['level'] == 'error' for i in issues):
        def available(items):
            return sum((Decimal(str(c['credits'])) for c in items), Decimal(0))
        limits = [('最低总学分', goals['total'], available(courses)),
                  ('最低选修学分', goals['elective'], available(c for c in courses if c['type'] == '选修'))]
        limits.extend((cat+'最低学分', target, available(c for c in courses if c['category'] == cat))
                      for cat, target in goals['categories'].items())
        for label, target, supply in limits:
            if target is not None and Decimal(str(target)) > supply:
                issues.append(dict(level='warning', index=0, id=courses[0]['id'], field='goals',
                                   message=f'{label} {target:g} 超过课程表可提供的 {supply:g} 分，请核对目标或补充课程'))
    selection = payload.get('selected', [])
    if not isinstance(selection, list) or any(not isinstance(x, str) for x in selection):
        raise ValueError('已选课程格式无效')
    selected = set(selection)
    selected.update(c['id'] for c in courses if c['type'] == '必修')
    plan = dict(name=name, termCount=count, goals=goals, courses=courses,
                selected=[c['id'] for c in courses if c['id'] in selected], source=txt(payload.get('source'))[:250])
    return plan, issues


def compute(plan):
    selected = set(plan['selected'])
    placed = [c for c in plan['courses'] if c['id'] in selected]
    def total(items):
        return float(sum((Decimal(str(c['credits'])) for c in items), Decimal(0)))
    categories = {cat: total(c for c in placed if c['category'] == cat) for cat in dict.fromkeys(c['category'] for c in plan['courses'])}
    checks = []
    goals = plan['goals']
    elective = total(c for c in placed if c['type'] == '选修')
    for label, actual, target in [('总学分', total(placed), goals['total']), ('选修学分', elective, goals['elective'])] + [(k, categories.get(k, 0), v) for k, v in goals['categories'].items()]:
        if target is not None:
            checks.append(dict(label=label, actual=actual, target=target, met=actual >= target, gap=round(max(0, target-actual), 2)))
    return dict(total=total(placed), required=total(c for c in placed if c['type'] == '必修'), elective=elective,
                selectedCount=len(placed), categories=categories, checks=checks,
                status='unknown' if not checks else 'met' if all(c['met'] for c in checks) else 'incomplete')
